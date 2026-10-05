import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT_NAME = "Arial"

HEADER_FILL = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name=FONT_NAME, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

POS_LABEL = {
    "noun": "существительное",
    "verb": "глагол",
    "adjective": "прилагательное",
    "adverb": "наречие",
    "phrase": "выражение",
    "other": "другое",
}

COLUMNS = [
    ("word", "Слово", 18),
    ("transcription", "Транскрипция", 16),
    ("part_of_speech", "Часть речи", 16),
    ("translation", "Перевод", 20),
    ("example_en", "Пример (EN)", 40),
    ("example_ru", "Пример (RU)", 40),
    ("synonyms", "Синонимы", 24),
    ("antonyms", "Антонимы", 24),
    ("learned", "Выучено", 12),
    ("created_at", "Добавлено", 18),
]


def _cell_value(card: dict, key: str):
    if key == "part_of_speech":
        return POS_LABEL.get(card.get("part_of_speech"), card.get("part_of_speech") or "")
    if key in ("synonyms", "antonyms"):
        items = card.get(key) or []
        return ", ".join(items)
    if key == "learned":
        return "Да" if card.get("learned") else "Нет"
    if key == "created_at":
        raw = card.get("created_at")
        if not raw:
            return ""
        try:
            dt = datetime.fromisoformat(raw.replace("Z", "+00:00"))
            return dt.strftime("%d.%m.%Y")
        except ValueError:
            return raw
    return card.get(key, "")


def build_workbook(cards: list[dict]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Слова"

    headers = [label for _, label, _ in COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER

    for card in cards:
        row = [_cell_value(card, key) for key, _, _ in COLUMNS]
        ws.append(row)

    last_row = len(cards) + 1
    for row_idx in range(2, last_row + 1):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER

    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    if cards:
        table_ref = f"A1:{get_column_letter(len(headers))}{last_row}"
        table = Table(displayName="Words", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showRowStripes=True,
            showFirstColumn=False,
            showLastColumn=False,
            showColumnStripes=False,
        )
        ws.add_table(table)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer